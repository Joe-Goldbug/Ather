#!/usr/bin/env python3
"""Generate the EVA investor pitch deck V2."""

from pathlib import Path
from typing import Iterable, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "pitch-deck" / "eva-investor-pitch-v2.pptx"

BG = "080A12"
PANEL = "111521"
PANEL_2 = "171C2B"
CYAN = "39F2E6"
PURPLE = "A779FF"
GOLD = "FFC857"
WHITE = "F5F7FB"
MUTED = "98A2B3"
RED = "FF6B7A"
GREEN = "6EE7B7"
FONT_CN = "PingFang SC"
FONT_EN = "Avenir Next"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "EVA Investor Pitch V2"
prs.core_properties.subject = "Pre-seed investor deck"
prs.core_properties.author = "EVA"
prs.core_properties.comments = "Generated from verified project facts on 2026-07-26."


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             size: float = 18, color: str = WHITE, bold: bool = False,
             font: str = FONT_CN, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.MIDDLE, margin: float = 0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x: float, y: float, w: float, h: float,
             fill: str = PANEL, line_color: Optional[str] = None,
             radius: bool = True, line_width: float = 1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color or fill)
    shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float,
             color: str = MUTED, width: float = 1.4, dash=None):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_circle(slide, x: float, y: float, d: float,
               fill: str = PANEL_2, line_color: str = CYAN, width: float = 1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(width)
    return shape


def add_chip(slide, label: str, x: float, y: float, w: float,
             color: str = CYAN):
    add_rect(slide, x, y, w, 0.34, fill=PANEL_2, line_color=color, line_width=0.8)
    add_text(slide, label, x, y, w, 0.34, size=10, color=color,
             bold=True, align=PP_ALIGN.CENTER)


def add_base(slide, number: int, section: str) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(BG)
    add_rect(slide, 0, 0, 13.333, 0.055, fill=CYAN, line_color=CYAN, radius=False)
    add_text(slide, "EVA / INVESTOR DECK", 0.62, 0.24, 2.7, 0.28,
             size=9, color=CYAN, bold=True, font=FONT_EN)
    add_text(slide, section.upper(), 10.1, 0.24, 2.0, 0.28,
             size=9, color=MUTED, bold=True, font=FONT_EN, align=PP_ALIGN.RIGHT)
    add_text(slide, f"{number:02d}", 12.25, 0.24, 0.44, 0.28,
             size=9, color=MUTED, bold=True, font=FONT_EN, align=PP_ALIGN.RIGHT)
    add_line(slide, 0.62, 7.12, 12.7, 7.12, color=PANEL_2, width=0.7)
    add_text(slide, "Evidence you can inspect. A model you can revise.",
             0.62, 7.16, 4.6, 0.18, size=7.5, color=MUTED, font=FONT_EN)


def add_title(slide, heading: str, sub: str = "") -> None:
    add_text(slide, heading, 0.62, 0.72, 11.9, 0.72,
             size=28, bold=True, valign=MSO_ANCHOR.BOTTOM)
    if sub:
        add_text(slide, sub, 0.64, 1.47, 11.7, 0.44,
                 size=13, color=MUTED, valign=MSO_ANCHOR.TOP)


def new_slide(number: int, section: str, heading: str = "", sub: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, number, section)
    if heading:
        add_title(slide, heading, sub)
    return slide


def add_arrow(slide, x: float, y: float, color: str = MUTED) -> None:
    add_text(slide, "→", x, y, 0.35, 0.4, size=20, color=color,
             bold=True, align=PP_ALIGN.CENTER)


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = rgb(BG)
add_rect(slide, 0, 0, 13.333, 0.055, fill=CYAN, line_color=CYAN, radius=False)
add_text(slide, "EVA", 0.7, 1.15, 5.6, 0.85, size=52, color=WHITE,
         bold=True, font=FONT_EN, valign=MSO_ANCHOR.BOTTOM)
add_text(slide, "一面会被你纠正的心智镜像", 0.72, 2.12, 6.0, 0.78,
         size=31, color=WHITE, bold=True)
add_text(slide, "Evidence you can inspect.\nA model you can revise.",
         0.75, 3.12, 4.8, 0.9, size=17, color=CYAN, bold=True, font=FONT_EN)
add_chip(slide, "PRE-SEED", 0.75, 4.42, 1.25, PURPLE)
add_text(slide, "Investor Pitch · 2026.07", 2.2, 4.41, 2.8, 0.36,
         size=11, color=MUTED, font=FONT_EN)
# Abstract evidence network
nodes = [(8.2, 1.0), (10.0, 0.86), (11.2, 1.62), (8.75, 2.45),
         (10.45, 2.42), (11.65, 3.0), (8.05, 3.85), (9.7, 4.2),
         (11.25, 4.5), (9.0, 5.55), (10.8, 5.92)]
links = [(0,1),(1,2),(0,3),(1,4),(2,5),(3,4),(4,5),(3,6),(4,7),
         (5,8),(6,7),(7,8),(6,9),(7,9),(7,10),(8,10),(9,10)]
for a, b in links:
    add_line(slide, nodes[a][0] + .12, nodes[a][1] + .12,
             nodes[b][0] + .12, nodes[b][1] + .12,
             color=PURPLE if (a + b) % 3 == 0 else "263047", width=1.2)
for i, (x, y) in enumerate(nodes):
    add_circle(slide, x, y, 0.24, fill=BG,
               line_color=CYAN if i in (1,4,7,10) else PURPLE, width=1.4)
add_text(slide, "USER-CORRECTABLE EVIDENCE MODEL", 7.95, 6.47, 4.2, 0.3,
         size=9, color=MUTED, font=FONT_EN, align=PP_ALIGN.RIGHT)


# Slide 2
slide = new_slide(2, "Problem", "今天的自我理解工具，在答案出现时就停止了")
panels = [
    ("01", "一次性测试", "给出标签\n却不追踪现实", RED),
    ("02", "通用 AI 聊天", "给出回应\n却混淆推测与事实", PURPLE),
    ("03", "共同缺口", "没有长期、可核查、\n可纠正的证据链", CYAN),
]
for i, (num, name, body, accent) in enumerate(panels):
    x = 0.72 + i * 4.18
    add_rect(slide, x, 2.05, 3.75, 3.75, fill=PANEL, line_color="222A3A")
    add_text(slide, num, x + 0.25, 2.28, 0.62, 0.42, size=12,
             color=accent, bold=True, font=FONT_EN)
    add_text(slide, name, x + 0.25, 2.85, 3.18, 0.55, size=22, bold=True)
    add_text(slide, body, x + 0.25, 3.55, 3.15, 1.05, size=17,
             color=MUTED, valign=MSO_ANCHOR.TOP)
    add_rect(slide, x + 0.25, 5.25, 1.0, 0.055, fill=accent,
             line_color=accent, radius=False)
add_text(slide, "依据是什么？    我不同意怎么办？    变化如何被发现？",
         1.2, 6.15, 10.9, 0.48, size=17, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)

# Slide 3
slide = new_slide(3, "Thesis", "自我理解不是一次判定，而是一条证据循环",
                  "画像是工作假设，不是人格判决。")
center_x, center_y = 5.55, 2.85
add_circle(slide, center_x, center_y, 2.15, fill=PANEL, line_color=CYAN, width=2)
add_text(slide, "可修订\n画像", center_x + 0.3, center_y + 0.5, 1.55, 1.1,
         size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
loop_items = [
    ("HYPOTHESIS", "有限假设", 1.0, 2.1, PURPLE),
    ("EVIDENCE", "现实证据", 9.75, 2.1, CYAN),
    ("CORRECTION", "用户纠正", 9.75, 4.65, GOLD),
    ("UPDATE", "置信更新", 1.0, 4.65, GREEN),
]
for en, cn, x, y, accent in loop_items:
    add_rect(slide, x, y, 2.55, 1.05, fill=PANEL, line_color=accent, line_width=1.4)
    add_text(slide, en, x + .16, y + .12, 2.2, .25, size=9,
             color=accent, bold=True, font=FONT_EN)
    add_text(slide, cn, x + .16, y + .39, 2.2, .48, size=17, bold=True)
add_line(slide, 3.55, 2.62, 5.5, 3.25, color=PURPLE, width=1.5)
add_line(slide, 7.7, 3.25, 9.75, 2.62, color=CYAN, width=1.5)
add_line(slide, 10.95, 3.16, 10.95, 4.65, color=GOLD, width=1.5)
add_line(slide, 9.75, 5.18, 7.7, 4.58, color=GOLD, width=1.5)
add_line(slide, 5.5, 4.58, 3.55, 5.18, color=GREEN, width=1.5)
add_line(slide, 2.25, 4.65, 2.25, 3.16, color=GREEN, width=1.5)
add_chip(slide, "AI 解释默认是候选假设", 4.65, 6.15, 4.05, PURPLE)

# Slide 4
slide = new_slide(4, "Product", "从 14 题开始，之后每次只做一个小动作",
                  "低摩擦输入 × 可追溯证据 × 持续更新")
steps = [
    ("14", "结构化基线", "CURRENT"), ("01", "初步画像", "CURRENT"),
    ("+", "现实记录", "CURRENT"), ("01", "单题校准", "CURRENT"),
    ("↺", "用户纠正", "CURRENT"), ("◫", "历史快照", "VERIFY"),
]
for i, (icon, label, status) in enumerate(steps):
    x = 0.52 + i * 2.13
    accent = GOLD if status == "VERIFY" else CYAN
    add_circle(slide, x + 0.36, 2.35, 1.12, fill=PANEL, line_color=accent, width=1.5)
    add_text(slide, icon, x + 0.50, 2.62, .84, .56, size=22, color=accent,
             bold=True, font=FONT_EN, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 3.68, 1.85, .45, size=15, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, status, x + .25, 4.24, 1.35, .25, size=8,
             color=accent, bold=True, font=FONT_EN, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_arrow(slide, x + 1.79, 2.73, color="4B556B")
add_rect(slide, 1.0, 5.1, 11.25, 0.9, fill=PANEL, line_color="253049")
add_text(slide, "一次建立基线    ·    日常补充证据    ·    在分歧处纠正    ·    随时间重新校准",
         1.25, 5.28, 10.75, .48, size=16, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(slide, "注：周回顾 / 快照端到端运行仍待核验", 8.8, 6.35, 3.4, .25,
         size=9, color=GOLD, align=PP_ALIGN.RIGHT)


# Slide 5
slide = new_slide(5, "Experience", "用户永远拥有最后解释权")
# Original words card
add_rect(slide, 0.8, 2.0, 3.55, 3.65, fill=PANEL, line_color="283148")
add_chip(slide, "01 / USER WORDS", 1.08, 2.25, 1.55, CYAN)
add_text(slide, "“我不是在回避冲突，\n我是在等彼此冷静。”", 1.08, 3.0, 2.95, 1.12,
         size=20, color=WHITE, bold=True, valign=MSO_ANCHOR.TOP)
add_text(slide, "原话保留 · 不被覆盖", 1.08, 4.82, 2.7, .35,
         size=11, color=MUTED)
# AI hypothesis
add_rect(slide, 4.9, 2.0, 3.55, 3.65, fill=PANEL, line_color=PURPLE)
add_chip(slide, "02 / AI HYPOTHESIS", 5.18, 2.25, 1.85, PURPLE)
add_text(slide, "你可能倾向于\n回避冲突。", 5.18, 3.0, 2.95, 1.12,
         size=20, color=WHITE, bold=True, valign=MSO_ANCHOR.TOP)
add_chip(slide, "PENDING / 待确认", 5.18, 4.78, 1.75, GOLD)
# User control
add_rect(slide, 9.0, 2.0, 3.55, 3.65, fill=PANEL, line_color=CYAN)
add_chip(slide, "03 / USER CONTROL", 9.28, 2.25, 1.75, CYAN)
add_text(slide, "不准确。\n让我告诉你真实原因。", 9.28, 3.0, 2.95, 1.12,
         size=20, color=WHITE, bold=True, valign=MSO_ANCHOR.TOP)
add_rect(slide, 9.28, 4.66, 1.35, .52, fill=CYAN, line_color=CYAN)
add_text(slide, "提交纠正", 9.28, 4.66, 1.35, .52, size=12, color=BG,
         bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "被说错之后，系统如何修正，才是信任发生的地方。",
         1.1, 6.2, 11.1, .42, size=17, color=CYAN, bold=True,
         align=PP_ALIGN.CENTER)

# Slide 6
slide = new_slide(6, "Beachhead", "从拒绝被一个标签概括的人开始",
                  "首批用户定义是待验证假设，不是已证明市场结论。")
add_rect(slide, 0.75, 2.1, 4.1, 3.85, fill=PANEL, line_color=PURPLE)
add_chip(slide, "EARLY USER HYPOTHESIS", 1.05, 2.42, 2.15, PURPLE)
add_text(slide, "已经在做自我探索，\n但不再满足于\n一次性标签与泛化建议。",
         1.05, 3.05, 3.45, 1.65, size=22, color=WHITE, bold=True,
         valign=MSO_ANCHOR.TOP)
add_text(slide, "愿意记录 · 愿意复盘 · 愿意纠正", 1.05, 5.18, 3.4, .35,
         size=11, color=MUTED)
use_cases = [
    ("关系", "看见重复出现的互动模式", CYAN),
    ("压力", "理解高压情境下的反应", GOLD),
    ("决定", "在重要选择后复盘自己", GREEN),
]
for i, (name, body, accent) in enumerate(use_cases):
    y = 2.1 + i * 1.3
    add_rect(slide, 5.45, y, 6.85, 1.02, fill=PANEL, line_color="283148")
    add_circle(slide, 5.73, y + .22, .56, fill=BG, line_color=accent, width=1.3)
    add_text(slide, name[0], 5.83, y + .31, .36, .33, size=13,
             color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, name, 6.55, y + .18, 1.05, .33, size=15, bold=True)
    add_text(slide, body, 7.68, y + .18, 4.15, .5, size=14, color=MUTED)
add_chip(slide, "NEXT: INTERVIEW → ACTIVATION → D7 RETURN", 6.28, 6.15, 5.1, CYAN)

# Slide 7
slide = new_slide(7, "Category", "从“给答案”走向“维护一个可修订模型”")
# Axes
x0, y0, x1, y1 = 1.55, 5.85, 11.9, 2.05
add_line(slide, x0, y0, x1, y0, color=MUTED, width=1.5)
add_line(slide, x0, y0, x0, y1, color=MUTED, width=1.5)
add_text(slide, "一次性", 1.32, 6.08, 1.0, .3, size=10, color=MUTED)
add_text(slide, "纵向持续", 10.9, 6.08, 1.2, .3, size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
add_text(slide, "黑盒结论", .62, 5.45, .75, .55, size=10, color=MUTED,
         align=PP_ALIGN.CENTER)
add_text(slide, "证据可见\n用户可纠正", .55, 1.95, .9, .72, size=10, color=MUTED,
         align=PP_ALIGN.CENTER)
# Grid
add_line(slide, 6.72, 2.15, 6.72, 5.85, color="20283A", width=1)
add_line(slide, 1.55, 3.95, 11.9, 3.95, color="20283A", width=1)
positions = [
    ("一次性人格测试", 2.25, 4.8, RED),
    ("AI 陪伴 / 日记", 8.0, 4.8, PURPLE),
    ("专业测量工具", 2.35, 2.8, GOLD),
]
for label, x, y, accent in positions:
    add_rect(slide, x, y, 2.25, .58, fill=PANEL, line_color=accent)
    add_text(slide, label, x, y, 2.25, .58, size=11, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
add_rect(slide, 8.4, 2.45, 2.55, 1.0, fill="102327", line_color=CYAN, line_width=2)
add_text(slide, "EVA", 8.4, 2.55, 2.55, .38, size=20, color=CYAN,
         bold=True, font=FONT_EN, align=PP_ALIGN.CENTER)
add_text(slide, "Consumer self-discovery × evidence loop", 8.4, 2.96, 2.55, .24,
         size=7.5, color=MUTED, font=FONT_EN, align=PP_ALIGN.CENTER)


# Slide 8
slide = new_slide(8, "Business", "基础自我理解免费，深度纵向洞察订阅")
plans = [
    ("FREE", "建立信任", ["一次基线与初步画像", "基础证据时间线", "查看置信度", "纠正 / 删除 / 导出"], CYAN),
    ("PRO", "放大长期价值", ["更多动态模拟", "深度趋势与跨月对比", "高级周报 / 月报", "语音与高级输入配额"], PURPLE),
]
for i, (plan, purpose, items, accent) in enumerate(plans):
    x = .85 + i * 6.1
    add_rect(slide, x, 2.0, 5.55, 3.95, fill=PANEL, line_color=accent, line_width=1.5)
    add_text(slide, plan, x + .35, 2.3, 1.25, .5, size=23, color=accent,
             bold=True, font=FONT_EN)
    add_text(slide, purpose, x + 2.65, 2.38, 2.4, .35, size=14,
             color=WHITE, bold=True, align=PP_ALIGN.RIGHT)
    for j, item in enumerate(items):
        yy = 3.05 + j * .58
        add_circle(slide, x + .38, yy + .09, .18, fill=accent,
                   line_color=accent, width=.5)
        add_text(slide, item, x + .75, yy, 4.35, .38, size=13.5, color=WHITE)
add_chip(slide, "NO PAYWALL FOR CORRECTION RIGHTS", 1.25, 6.28, 3.55, GOLD)
add_text(slide, "不按脆弱时刻收费 · 不利用焦虑变现 · 定价待用户验证",
         5.2, 6.27, 6.8, .36, size=12, color=MUTED, align=PP_ALIGN.RIGHT)

# Slide 9
slide = new_slide(9, "Moat", "不是聊天记录，而是可审计的长期证据关系",
                  "Evidence Graph + Revision History + Confidence Engine")
# Evidence nodes left
left_nodes = [(1.0,2.35,"选择",CYAN),(2.35,2.0,"记录",PURPLE),(3.35,3.0,"纠正",GOLD),
              (1.5,4.0,"决定",GREEN),(2.75,4.7,"确认",CYAN)]
for i, (x,y,label,accent) in enumerate(left_nodes):
    if i:
        add_line(slide, 2.2, 3.45, x+.42, y+.42, color="33405A", width=1.1)
    add_circle(slide, x, y, .84, fill=PANEL, line_color=accent, width=1.4)
    add_text(slide, label, x+.08, y+.22, .68, .38, size=11, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "长期证据图谱", .9, 5.65, 3.3, .34, size=13, color=MUTED,
         align=PP_ALIGN.CENTER)
add_arrow(slide, 4.45, 3.45, CYAN)
# Four factors
factors = [("充分度",CYAN),("一致性",PURPLE),("来源覆盖",GOLD),("校准度",GREEN)]
for i,(name,accent) in enumerate(factors):
    x = 5.0 + (i%2)*1.75
    y = 2.25 + (i//2)*1.35
    add_rect(slide, x, y, 1.5, .92, fill=PANEL, line_color=accent)
    add_text(slide, name, x, y+.18, 1.5, .45, size=13, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
add_rect(slide, 5.0, 5.05, 3.25, .74, fill="102327", line_color=CYAN)
add_text(slide, "CONFIDENCE = MIN(4 FACTORS)", 5.0, 5.05, 3.25, .74,
         size=11, color=CYAN, bold=True, font=FONT_EN, align=PP_ALIGN.CENTER)
add_arrow(slide, 8.55, 3.45, CYAN)
# Outcome
add_rect(slide, 9.1, 2.18, 3.25, 3.5, fill=PANEL, line_color=CYAN, line_width=1.6)
add_chip(slide, "REVISABLE PORTRAIT", 9.45, 2.48, 2.15, CYAN)
add_text(slide, "有来源\n有边界\n有修订历史", 9.52, 3.22, 2.4, 1.45,
         size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "数据关系更可信，而不是数据越多。", 9.35, 5.12, 2.75, .36,
         size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

# Slide 10
slide = new_slide(10, "Validation", "先证明价值与可信度，再追求规模",
                  "V1 Evidence Loop — In Progress")
gates = [
    ("01", "价值", "用户是否愿意回来\n补充现实证据？", CYAN),
    ("02", "可信", "用户是否理解边界\n并愿意纠正？", PURPLE),
    ("03", "留存", "纵向回顾是否形成\n重复使用理由？", GOLD),
]
for i,(num,name,question,accent) in enumerate(gates):
    x = .78 + i*4.18
    add_rect(slide, x, 2.05, 3.72, 3.55, fill=PANEL, line_color=accent,
             line_width=1.4 if i < 2 else .8)
    add_text(slide, num, x+.28, 2.32, .65, .4, size=12, color=accent,
             bold=True, font=FONT_EN)
    add_text(slide, name, x+.28, 2.9, 2.9, .55, size=23, bold=True)
    add_text(slide, question, x+.28, 3.72, 3.0, 1.0, size=16, color=MUTED,
             valign=MSO_ANCHOR.TOP)
    status = "TESTING" if i == 0 else ("DESIGNED" if i == 1 else "PENDING")
    add_chip(slide, status, x+.28, 5.02, 1.1, accent)
add_text(slide, "当前：主链路已定义  ·  纵向变化检测未交付  ·  商业指标待采集",
         1.0, 6.18, 11.2, .42, size=13, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)


# Slide 11
slide = new_slide(11, "Team", "为什么是我们",
                  "正式路演前请替换全部占位符，只放可核验经历与结果。")
team = [
    ("【创始人姓名】", "PRODUCT / ENGINEERING", "一句话写清与长期自我模型\n直接相关的独特经历", CYAN),
    ("【核心成员】", "SCIENCE / DATA", "写明心理测量、研究或数据能力\n以及可验证成果", PURPLE),
    ("【核心成员】", "GROWTH / DESIGN", "写明消费产品、订阅或设计能力\n以及可验证成果", GOLD),
]
for i,(name,role,bio,accent) in enumerate(team):
    x = .72 + i*4.18
    add_rect(slide, x, 2.0, 3.72, 3.62, fill=PANEL, line_color="283148")
    add_circle(slide, x+.27, 2.3, .9, fill=BG, line_color=accent, width=1.5)
    add_text(slide, "＋", x+.43, 2.49, .58, .38, size=19, color=accent,
             align=PP_ALIGN.CENTER)
    add_text(slide, name, x+.27, 3.45, 3.1, .48, size=20, bold=True)
    add_text(slide, role, x+.27, 3.97, 3.1, .3, size=9, color=accent,
             bold=True, font=FONT_EN)
    add_text(slide, bio, x+.27, 4.5, 3.05, .78, size=12.5, color=MUTED,
             valign=MSO_ANCHOR.TOP)
add_rect(slide, 1.15, 6.05, 11.0, .58, fill=PANEL_2, line_color="303A52")
add_text(slide, "正在寻找：心理测量顾问  ·  隐私 / 安全顾问  ·  种子用户社群伙伴",
         1.15, 6.05, 11.0, .58, size=13, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)

# Slide 12
slide = new_slide(12, "Ask", "共同完成“可信的长期自我模型”第一次验证")
add_text(slide, "WE ARE RAISING", .78, 1.72, 2.2, .35, size=10,
         color=CYAN, bold=True, font=FONT_EN)
add_text(slide, "【融资金额】", .78, 2.05, 4.15, .92, size=42,
         color=WHITE, bold=True)
add_text(slide, "Pre-seed · 【runway】个月 · 【出让比例】", .8, 3.0, 4.5, .4,
         size=13, color=MUTED, font=FONT_EN)
asks = [
    ("资本", "完成 V1 闭环、用户验证与科学验证设计", CYAN),
    ("试点", "100–300 名愿意持续 4–8 周的种子用户", PURPLE),
    ("能力", "心理测量、消费订阅与隐私安全伙伴", GOLD),
]
for i,(name,body,accent) in enumerate(asks):
    y = 1.82 + i*1.3
    add_rect(slide, 6.15, y, 6.15, 1.0, fill=PANEL, line_color="283148")
    add_chip(slide, name, 6.42, y+.32, .82, accent)
    add_text(slide, body, 7.55, y+.2, 4.35, .56, size=13.5, color=WHITE,
             bold=True)
add_rect(slide, .78, 4.22, 4.65, 1.2, fill=PANEL, line_color=CYAN)
add_text(slide, "【产品二维码】", 1.03, 4.47, 1.2, .58, size=11,
         color=MUTED, align=PP_ALIGN.CENTER)
add_text(slide, "【联系人 / 邮箱】", 2.47, 4.58, 2.45, .35,
         size=13, color=WHITE, bold=True)
add_text(slide, "AI 不该替你定义你。\n它应该帮助你看见证据，\n并保留改写自己的权利。",
         .8, 5.68, 11.55, .95, size=24, color=CYAN, bold=True,
         align=PP_ALIGN.CENTER)

# Save
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Generated {len(prs.slides)} slides: {OUTPUT}")
